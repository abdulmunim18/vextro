"""Bulk CSV/XLSX import for SME-managed products."""

from __future__ import annotations

import csv
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO
from pathlib import Path

from fastapi import HTTPException, status
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy.orm import Session

from app.schemas.sme import (
    BusinessProductCreate,
    BusinessProductImportError,
    BusinessProductImportResponse,
)
from app.services.sme_service import SMEService


class BusinessProductImportService:
    """Validate and create business products from a spreadsheet."""

    ALLOWED_EXTENSIONS = {".csv", ".xlsx"}
    MAX_FILE_SIZE_BYTES = 2 * 1024 * 1024
    MAX_ROWS = 5000
    REQUIRED_COLUMNS = {"name", "sku"}
    ALLOWED_COLUMNS = {
        "name",
        "sku",
        "catalog_product_id",
        "catalog_product_name",
        "cost_price",
        "selling_price",
        "currency",
        "stock_level",
        "reorder_level",
    }

    def __init__(self, sme_service: SMEService | None = None) -> None:
        self.sme_service = sme_service or SMEService()

    @staticmethod
    def _clean(value: object) -> str:
        if value is None:
            return ""
        return str(value).strip()

    @classmethod
    def _normalize_row(cls, row: dict[object, object]) -> dict[str, str]:
        return {
            cls._clean(key).casefold(): cls._clean(value)
            for key, value in row.items()
            if cls._clean(key)
        }

    def _read_rows(
        self,
        *,
        filename: str,
        file_content: bytes,
    ) -> list[tuple[int, dict[str, str]]]:
        extension = Path(filename).suffix.casefold()
        if extension not in self.ALLOWED_EXTENSIONS:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail="Upload a .csv or .xlsx product file.",
            )
        if not file_content:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail="The uploaded product file is empty.",
            )
        if len(file_content) > self.MAX_FILE_SIZE_BYTES:
            raise HTTPException(
                status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
                detail="The product import file must not exceed 2 MB.",
            )

        rows: list[tuple[int, dict[str, str]]] = []
        if extension == ".csv":
            try:
                decoded = file_content.decode("utf-8-sig")
            except UnicodeDecodeError as error:
                raise HTTPException(
                    status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                    detail="CSV files must use UTF-8 encoding.",
                ) from error
            reader = csv.DictReader(StringIO(decoded, newline=""), strict=True)
            headers = [self._clean(item).casefold() for item in reader.fieldnames or []]
            for row_number, row in enumerate(reader, start=2):
                normalized = self._normalize_row(row)
                if any(normalized.values()):
                    rows.append((row_number, normalized))
        else:
            try:
                workbook = load_workbook(BytesIO(file_content), read_only=True, data_only=True)
                sheet = workbook.active
                values = sheet.iter_rows(values_only=True)
                header_row = next(values, None)
                headers = [self._clean(item).casefold() for item in header_row or []]
                for row_number, values_row in enumerate(values, start=2):
                    normalized = {
                        headers[index]: self._clean(value)
                        for index, value in enumerate(values_row)
                        if index < len(headers) and headers[index]
                    }
                    if any(normalized.values()):
                        rows.append((row_number, normalized))
            except (OSError, ValueError, KeyError) as error:
                raise HTTPException(
                    status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                    detail="The Excel product file could not be read.",
                ) from error
            finally:
                if "workbook" in locals():
                    workbook.close()

        missing = sorted(self.REQUIRED_COLUMNS - set(headers))
        unknown = sorted(set(headers) - self.ALLOWED_COLUMNS)
        if missing:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail=f"Missing required columns: {', '.join(missing)}.",
            )
        if unknown:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail=f"Unsupported columns: {', '.join(unknown)}.",
            )
        if len(rows) > self.MAX_ROWS:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail="A product import may contain at most 5,000 rows.",
            )
        return rows

    @staticmethod
    def _decimal(value: str) -> Decimal | None:
        if not value:
            return None
        try:
            return Decimal(value.replace(",", ""))
        except InvalidOperation as error:
            raise ValueError(f"'{value}' is not a valid price.") from error

    @staticmethod
    def _integer(value: str, *, default: int = 0) -> int:
        if not value:
            return default
        try:
            parsed = Decimal(value.replace(",", ""))
            if parsed != parsed.to_integral_value():
                raise ValueError
            return int(parsed)
        except (InvalidOperation, ValueError) as error:
            raise ValueError(f"'{value}' is not a valid whole number.") from error

    @staticmethod
    def _error_message(error: Exception) -> str:
        if isinstance(error, HTTPException):
            detail = error.detail
            if isinstance(detail, dict):
                return str(detail.get("message") or detail.get("detail") or detail)
            return str(detail)
        if isinstance(error, ValidationError):
            first = error.errors()[0]
            return str(first.get("msg") or "Invalid product data.")
        return str(error) or "The product row could not be imported."

    def import_products(
        self,
        database_session: Session,
        *,
        organization_id: int,
        user_id: int,
        filename: str | None,
        file_content: bytes,
    ) -> BusinessProductImportResponse:
        """Import valid rows while returning bounded errors for bad rows."""

        safe_filename = Path(filename or "products.csv").name[:255]
        rows = self._read_rows(filename=safe_filename, file_content=file_content)
        products = []
        errors = []
        catalog_matched_rows = 0

        for row_number, row in rows:
            sku = row.get("sku") or None
            try:
                if not row.get("name"):
                    raise ValueError("Product name is required.")
                if not sku:
                    raise ValueError("SKU is required for product imports.")

                catalog_product_id = None
                raw_catalog_id = row.get("catalog_product_id", "")
                if raw_catalog_id:
                    catalog_product_id = self._integer(raw_catalog_id)
                    if catalog_product_id < 1:
                        raise ValueError("Catalog product ID must be positive.")
                elif row.get("catalog_product_name"):
                    match = self.sme_service._resolve_catalog_product(
                        database_session,
                        row["catalog_product_name"],
                    )
                    catalog_product_id = match.id if match is not None else None

                payload = BusinessProductCreate(
                    canonical_product_id=catalog_product_id,
                    name=row["name"],
                    sku=sku,
                    cost_price=self._decimal(row.get("cost_price", "")),
                    selling_price=self._decimal(row.get("selling_price", "")),
                    currency=row.get("currency") or "PKR",
                    stock_level=self._integer(row.get("stock_level", "")),
                    reorder_level=self._integer(row.get("reorder_level", "")),
                )
                created = self.sme_service.create_business_product(
                    database_session,
                    organization_id=organization_id,
                    user_id=user_id,
                    payload=payload,
                    catalog_reference_name=(
                        row.get("catalog_product_name")
                        or row["name"]
                    ),
                )
                products.append(created)
                if created.canonical_product_id is not None:
                    catalog_matched_rows += 1
            except (HTTPException, ValidationError, ValueError) as error:
                errors.append(
                    BusinessProductImportError(
                        row_number=row_number,
                        sku=sku,
                        message=self._error_message(error),
                    )
                )

        return BusinessProductImportResponse(
            filename=safe_filename,
            total_rows=len(rows),
            created_rows=len(products),
            rejected_rows=len(errors),
            catalog_matched_rows=catalog_matched_rows,
            unmatched_rows=len(products) - catalog_matched_rows,
            products=products,
            errors=errors[:100],
        )
